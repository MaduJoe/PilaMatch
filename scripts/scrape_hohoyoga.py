"""
호호요가 구인게시판 파싱 스크립트
- 필라테스 구인 [대강 서울] + [정규 서울] 각 10페이지
- 목록 → 상세 페이지 → 테이블 + 본문에서 연락처/이메일/주소/급여 추출
- Excel(.xlsx) 시트 분리 저장 (대강 / 정규)
"""

import asyncio
import re
import os
from datetime import datetime
from playwright.async_api import async_playwright
from openpyxl import Workbook


# ─── Config ───
HOHO_ID = "jaekeunv"
HOHO_PW = "100djrqjsek!"
LOGIN_URL = "https://www.hohoyoga.com/"
MAX_PAGES = 10

BOARDS = [
    {
        "name": "대강",
        "url": "https://www.hohoyoga.com/job_pilates_sub_seoul",
    },
    {
        "name": "정규",
        "url": "https://www.hohoyoga.com/job_pilates_seoul",
    },
]

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "data")
OUTPUT_FILE = os.path.join(
    OUTPUT_DIR,
    f"hohoyoga_jobs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
)

FIELDNAMES = [
    "제목", "모집여부", "업체명", "날짜_및_시간", "지역",
    "프로필_필수여부", "작성자", "작성일",
    "이메일", "전화번호", "주소", "급여", "본문_요약", "url",
]

# ─── Regex ───
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"01[016789][\-\s.]?\d{3,4}[\-\s.]?\d{4}")
PAY_RE = re.compile(
    r"(?:시급|급여|페이|수당|금액|그룹|개인|1:1|개인레슨)?\s*[/:]?\s*"
    r"(\d{1,3}(?:,\d{3})*\s*원|\d+만\s*원)",
    re.IGNORECASE,
)


async def login(page) -> bool:
    """호호요가 XE CMS AJAX 로그인"""
    print("[1] 로그인 중...")
    await page.goto(LOGIN_URL, wait_until="domcontentloaded")

    login_resp = await page.evaluate(
        """async ({ uid, upw }) => {
            const fd = new FormData();
            fd.append('user_id', uid);
            fd.append('password', upw);
            fd.append('act', 'procMemberLogin');
            fd.append('xe_validator_id', '');
            const r = await fetch('/index.php', { method: 'POST', body: fd, credentials: 'same-origin' });
            return r.status;
        }""",
        {"uid": HOHO_ID, "upw": HOHO_PW},
    )
    print(f"   응답: {login_resp}")

    await page.goto(BOARDS[0]["url"], wait_until="domcontentloaded")
    if await page.query_selector('a[href*="search_target=member_srl"]'):
        print("   ✓ 로그인 성공")
        return True
    print("   ✗ 로그인 실패")
    return False


async def get_post_links_page(page, board_url: str, page_num: int) -> list[dict]:
    """게시판 특정 페이지에서 글 목록 추출"""
    url = board_url if page_num == 1 else f"{board_url}?page={page_num}"
    await page.goto(url, wait_until="domcontentloaded")

    posts = []
    headers = await page.query_selector_all("table.bd_lst thead th")
    col_count = len(headers)
    rows = await page.query_selector_all("table.bd_lst tbody tr")

    for row in rows:
        cells = await row.query_selector_all("td")
        if len(cells) < 5:
            continue

        if col_count >= 6:
            모집여부 = (await cells[1].inner_text()).strip()
            날짜시간 = (await cells[2].inner_text()).strip()
            지역 = (await cells[3].inner_text()).strip()
            title_cell = cells[4]
            글쓴이 = (await cells[5].inner_text()).strip() if len(cells) > 5 else ""
        else:
            모집여부 = (await cells[0].inner_text()).strip()
            날짜시간 = (await cells[1].inner_text()).strip()
            지역 = (await cells[2].inner_text()).strip()
            title_cell = cells[3]
            글쓴이 = (await cells[4].inner_text()).strip()

        link_el = await title_cell.query_selector("a")
        if not link_el:
            continue
        title = (await link_el.inner_text()).strip()
        href = await link_el.get_attribute("href")
        if not href:
            continue
        if not href.startswith("http"):
            href = "https://www.hohoyoga.com" + href

        posts.append({
            "title": title, "url": href,
            "모집여부": 모집여부, "날짜_및_시간": 날짜시간,
            "지역": 지역, "글쓴이": 글쓴이,
        })

    return posts


async def parse_detail(page, url: str) -> dict:
    """상세 페이지에서 테이블 + 본문 파싱"""
    await page.goto(url, wait_until="domcontentloaded")

    result = {
        "url": url, "업체명": "", "프로필_필수여부": "",
        "작성일": "", "이메일": "", "전화번호": "",
        "주소": "", "급여": "", "본문_요약": "",
    }

    for tr in await page.query_selector_all("table.et_vars tr"):
        th = await tr.query_selector("th")
        td = await tr.query_selector("td")
        if th and td:
            key = (await th.inner_text()).strip()
            val = (await td.inner_text()).strip()
            if "업체" in key:
                result["업체명"] = val
            elif "프로필" in key:
                result["프로필_필수여부"] = val

    date_el = await page.query_selector(".date")
    if date_el:
        result["작성일"] = (await date_el.inner_text()).strip()

    article = await page.query_selector("article .xe_content, article")
    if article:
        body_text = (await article.inner_text()).strip()
        emails = EMAIL_RE.findall(body_text)
        result["이메일"] = ", ".join(set(emails)) if emails else ""
        phones = PHONE_RE.findall(body_text)
        result["전화번호"] = ", ".join(set(phones)) if phones else ""
        addr_match = re.search(r"(?:위치|주소)\s*[:：]\s*(.+?)(?:\n|\(|$)", body_text)
        if addr_match:
            result["주소"] = addr_match.group(1).strip()
        pays = PAY_RE.findall(body_text)
        result["급여"] = ", ".join(set(pays)) if pays else ""
        clean = re.sub(r"\s+", " ", body_text)
        result["본문_요약"] = clean[:200]

    return result


async def scrape_board(context, board: dict) -> list[dict]:
    """게시판 10페이지 파싱"""
    list_page = await context.new_page()
    all_posts = []

    for pg in range(1, MAX_PAGES + 1):
        print(f"   [{board['name']}] 페이지 {pg}/{MAX_PAGES} 목록 수집...")
        posts = await get_post_links_page(list_page, board["url"], pg)
        if not posts:
            print(f"   [{board['name']}] 페이지 {pg}: 글 없음 → 종료")
            break
        all_posts.extend(posts)

    await list_page.close()
    print(f"   [{board['name']}] 총 {len(all_posts)}개 목록 수집 완료, 상세 파싱 시작...")

    results = []
    sem = asyncio.Semaphore(5)  # 동시 5개 탭

    async def fetch_detail(idx, post):
        async with sem:
            detail_page = await context.new_page()
            try:
                detail = await parse_detail(detail_page, post["url"])
                detail["제목"] = post["title"]
                detail["모집여부"] = post["모집여부"]
                detail["날짜_및_시간"] = post["날짜_및_시간"]
                detail["지역"] = post["지역"]
                detail["작성자"] = post["글쓴이"]
                if (idx + 1) % 20 == 0:
                    print(f"   [{board['name']}] 상세 {idx+1}/{len(all_posts)} 완료")
                return detail
            except Exception as e:
                print(f"   ⚠ 파싱 실패: {post['url']} - {e}")
                return None
            finally:
                await detail_page.close()

    tasks = [fetch_detail(i, p) for i, p in enumerate(all_posts)]
    details = await asyncio.gather(*tasks)
    results = [d for d in details if d is not None]

    print(f"   ✓ [{board['name']}] {len(results)}개 상세 파싱 완료")
    return results


def save_xlsx(board_results: dict[str, list[dict]]):
    """Excel 파일로 시트 분리 저장"""
    wb = Workbook()
    wb.remove(wb.active)  # 기본 시트 제거

    for sheet_name, rows in board_results.items():
        ws = wb.create_sheet(title=sheet_name)
        ws.append(FIELDNAMES)
        for row in rows:
            ws.append([row.get(f, "") for f in FIELDNAMES])

        # 헤더 볼드 + 자동 너비
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)
        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    wb.save(OUTPUT_FILE)


async def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(
            headless=True,
            args=["--disable-gpu", "--disable-software-rasterizer", "--disable-dev-shm-usage"],
        )
        context = await browser.new_context(
            viewport={"width": 1280, "height": 720},
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
        )

        login_page = await context.new_page()
        if not await login(login_page):
            print("로그인 실패 - 중단")
            await browser.close()
            return
        await login_page.close()

        board_results = {}
        for board in BOARDS:
            results = await scrape_board(context, board)
            board_results[board["name"]] = results

        await browser.close()

    total = sum(len(v) for v in board_results.values())
    if not total:
        print("파싱 결과 없음")
        return

    save_xlsx(board_results)
    print(f"\n저장 완료: {OUTPUT_FILE}")
    for name, rows in board_results.items():
        print(f"  {name}: {len(rows)}개")
    print(f"  합계: {total}개")


if __name__ == "__main__":
    asyncio.run(main())
